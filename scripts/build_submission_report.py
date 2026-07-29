#!/usr/bin/env python3
"""Build the formal DOCX submission report from its Markdown authority."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = REPOSITORY_ROOT / "docs" / "submission-report.md"
DEFAULT_OUTPUT = (
    REPOSITORY_ROOT
    / "artifacts"
    / "submission_v2"
    / "report"
    / "草莓采摘URP项目总结报告.docx"
)
FIELD_OVERVIEW = (
    REPOSITORY_ROOT
    / "artifacts"
    / "submission_v2"
    / "figures"
    / "field_v3_overview.png"
)
WRIST_IMAGE = (
    REPOSITORY_ROOT
    / "results"
    / "development"
    / "field_v3_perception_shadow_v6"
    / "wrist_rgb.png"
)

BODY_FONT = "SimSun"
HEADING_FONT = "Microsoft YaHei"
LATIN_FONT = "Calibri"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "666666"
TABLE_FILL = "F4F6F9"
TABLE_ALT = "FAFBFC"
TABLE_BORDER = "AEB8C4"
WIDTH_DXA = 9360


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _set_run_font(
    run,
    *,
    font: str = BODY_FONT,
    latin_font: str = LATIN_FONT,
    size: float | None = None,
    bold: bool | None = None,
    color: str | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = latin_font
    run._element.get_or_add_rPr()
    fonts = run._element.rPr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin_font)
    fonts.set(qn("w:hAnsi"), latin_font)
    fonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def _configure_style(
    style,
    *,
    font: str,
    latin_font: str,
    size: float,
    color: str,
    before: float,
    after: float,
    line_spacing: float,
    bold: bool = False,
    keep_with_next: bool = False,
) -> None:
    style.font.name = latin_font
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin_font)
    fonts.set(qn("w:hAnsi"), latin_font)
    fonts.set(qn("w:eastAsia"), font)
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    fmt.keep_with_next = keep_with_next


def _set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (
        ("top", top),
        ("left", start),
        ("bottom", bottom),
        ("right", end),
    ):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)


def _set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def _table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != WIDTH_DXA:
        raise ValueError(f"table widths must sum to {WIDTH_DXA}: {widths}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_look = tbl_pr.find(qn("w:tblLook"))
    if tbl_look is not None:
        tbl_pr.remove(tbl_look)

    def insert_table_property(element) -> None:
        for tag in ("w:tblLayout", "w:tblLook"):
            successor = tbl_pr.find(qn(tag))
            if successor is not None:
                tbl_pr.insert(tbl_pr.index(successor), element)
                return
        tbl_pr.append(element)

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        insert_table_property(tbl_w)
    tbl_w.set(qn("w:w"), str(WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        insert_table_property(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        insert_table_property(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), TABLE_BORDER)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.insert(0, OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths[index])
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _column_widths(headers: list[str]) -> list[int]:
    count = len(headers)
    joined = "|".join(headers)
    if count == 4 and "SHA-256" in joined:
        return [1100, 1500, 1500, 5260]
    if count == 4:
        return [2550, 2100, 1750, 2960]
    if count == 3:
        return [2550, 4550, 2260]
    if count == 2:
        return [2800, 6560]
    return [WIDTH_DXA // count] * (count - 1) + [
        WIDTH_DXA - (WIDTH_DXA // count) * (count - 1)
    ]


def _clean_inline(value: str) -> str:
    return value.replace("`", "").replace("**", "").strip()


def _numbering(document: Document) -> tuple[int, int]:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(value.get(qn("w:abstractNumId")))
        for value in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(value.get(qn("w:numId")))
        for value in numbering.findall(qn("w:num"))
    ]

    def create(kind: str, start_id: int, num_id: int) -> int:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(start_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
        level.append(num_fmt)
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "space")
        level.append(suffix)
        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
        level.append(level_text)
        paragraph_props = OxmlElement("w:pPr")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "290")
        spacing.set(qn("w:lineRule"), "auto")
        paragraph_props.append(spacing)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "540")
        indent.set(qn("w:hanging"), "280")
        paragraph_props.append(indent)
        level.append(paragraph_props)
        run_props = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), LATIN_FONT)
        fonts.set(qn("w:hAnsi"), LATIN_FONT)
        fonts.set(qn("w:eastAsia"), BODY_FONT)
        run_props.append(fonts)
        level.append(run_props)
        abstract.append(level)
        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            numbering.insert(numbering.index(first_num), abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(start_id))
        num.append(abstract_ref)
        numbering.append(num)
        return num_id

    next_abstract = max(abstract_ids, default=0) + 1
    next_num = max(num_ids, default=0) + 1
    bullet = create("bullet", next_abstract, next_num)
    decimal = create("decimal", next_abstract + 1, next_num + 1)
    return bullet, decimal


def _clone_numbering_instance(document: Document, source_num_id: int) -> int:
    numbering = document.part.numbering_part.element
    source = next(
        (
            value
            for value in numbering.findall(qn("w:num"))
            if int(value.get(qn("w:numId"))) == source_num_id
        ),
        None,
    )
    if source is None:
        raise ValueError(f"numbering instance not found: {source_num_id}")
    abstract_ref = source.find(qn("w:abstractNumId"))
    if abstract_ref is None:
        raise ValueError(f"numbering instance has no abstract reference: {source_num_id}")
    current_ids = [
        int(value.get(qn("w:numId")))
        for value in numbering.findall(qn("w:num"))
    ]
    next_num_id = max(current_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(next_num_id))
    cloned_ref = OxmlElement("w:abstractNumId")
    cloned_ref.set(qn("w:val"), abstract_ref.get(qn("w:val")))
    num.append(cloned_ref)
    numbering.append(num)
    return next_num_id


def _set_list_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.insert(0, num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    value = OxmlElement("w:numId")
    value.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(value)


def _add_caption(document: Document, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(value)
    _set_run_font(run, size=9, color=MUTED)


def _add_figure(document: Document, path: Path, caption: str, width: float) -> None:
    if not path.is_file():
        raise ValueError(f"required report figure is missing: {path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    _add_caption(document, caption)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("第 ")
    _set_run_font(prefix, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    run_props.extend([color, size])
    run.extend([run_props, begin, instruction, separate, display, end])
    paragraph._p.append(run)
    suffix = paragraph.add_run(" 页")
    _set_run_font(suffix, size=9, color=MUTED)


def _cover(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(40)
    paragraph.paragraph_format.space_after = Pt(18)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("中国农业大学本科生 URP")
    _set_run_font(
        run,
        font=HEADING_FONT,
        size=12,
        bold=True,
        color=BLUE,
    )

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title_run = title.add_run(
        "基于 ROS 2、YOLO 与 Gazebo 的\n"
        "草莓成熟度识别及机械臂抓取搬运仿真"
    )
    _set_run_font(
        title_run,
        font=HEADING_FONT,
        size=24,
        bold=True,
        color=DARK_BLUE,
    )

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    subtitle_run = subtitle.add_run("提交版项目总结报告")
    _set_run_font(
        subtitle_run,
        font=HEADING_FONT,
        size=14,
        bold=True,
        color=MUTED,
    )

    _add_figure(
        document,
        FIELD_OVERVIEW,
        "field-v3 草莓田场景总览（用户模型的 Blender 渲染）",
        5.9,
    )

    meta = (
        "项目成员：刘泽林、韦丁元、张宇珊、罗泳、贾宇轩\n"
        "所在学院：工学院\n"
        "项目周期：2026 年 3 月 1 日—2027 年 3 月 1 日\n"
        "提交版日期：2026 年 7 月 29 日"
    )
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.35
    run = paragraph.add_run(meta)
    _set_run_font(run, size=10.5, color="333333")
    document.add_page_break()


def _parse_blocks(lines: list[str]) -> Iterable[tuple[str, object]]:
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("```"):
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index].rstrip())
                index += 1
            index += 1
            yield "code", "\n".join(code)
            continue
        if line.startswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                row = [
                    _clean_inline(value)
                    for value in lines[index].strip().strip("|").split("|")
                ]
                rows.append(row)
                index += 1
            if len(rows) >= 2 and all(
                re.fullmatch(r":?-{3,}:?", value.replace(" ", ""))
                for value in rows[1]
            ):
                rows.pop(1)
            yield "table", rows
            continue
        if line.startswith("### "):
            yield "h2", line[4:].strip()
            index += 1
            continue
        if line.startswith("## "):
            yield "h1", line[3:].strip()
            index += 1
            continue
        if line.startswith("- "):
            item = [line[2:]]
            index += 1
            while index < len(lines):
                next_line = lines[index].rstrip()
                if (
                    not next_line
                    or next_line.startswith(("#", "|", "-", "```"))
                    or re.match(r"^\d+\.\s+", next_line)
                ):
                    break
                item.append(next_line.strip())
                index += 1
            yield "bullet", _clean_inline(" ".join(item))
            continue
        match = re.match(r"^\d+\.\s+(.*)$", line)
        if match:
            item = [match.group(1)]
            index += 1
            while index < len(lines):
                next_line = lines[index].rstrip()
                if (
                    not next_line
                    or next_line.startswith(("#", "|", "-", "```"))
                    or re.match(r"^\d+\.\s+", next_line)
                ):
                    break
                item.append(next_line.strip())
                index += 1
            yield "number", _clean_inline(" ".join(item))
            continue
        paragraph = [line]
        index += 1
        while index < len(lines):
            next_line = lines[index].rstrip()
            if (
                not next_line
                or next_line.startswith(("#", "|", "-", "```"))
                or re.match(r"^\d+\.\s+", next_line)
            ):
                break
            paragraph.append(next_line)
            index += 1
        yield "paragraph", _clean_inline(" ".join(paragraph))


def build(source: Path, output: Path) -> dict[str, object]:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    _configure_style(
        normal,
        font=BODY_FONT,
        latin_font=LATIN_FONT,
        size=11,
        color="111111",
        before=0,
        after=8,
        line_spacing=1.333,
    )
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _configure_style(
        document.styles["Heading 1"],
        font=HEADING_FONT,
        latin_font=LATIN_FONT,
        size=16,
        color=BLUE,
        before=18,
        after=10,
        line_spacing=1.0,
        bold=True,
        keep_with_next=True,
    )
    _configure_style(
        document.styles["Heading 2"],
        font=HEADING_FONT,
        latin_font=LATIN_FONT,
        size=13,
        color=BLUE,
        before=12,
        after=6,
        line_spacing=1.0,
        bold=True,
        keep_with_next=True,
    )
    _configure_style(
        document.styles["Heading 3"],
        font=HEADING_FONT,
        latin_font=LATIN_FONT,
        size=12,
        color=DARK_BLUE,
        before=8,
        after=4,
        line_spacing=1.0,
        bold=True,
        keep_with_next=True,
    )

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("中国农业大学本科生 URP｜草莓采摘仿真项目总结报告")
    _set_run_font(run, size=9, color=MUTED)
    _add_page_number(section.footer.paragraphs[0])

    bullet_num_id, decimal_num_id = _numbering(document)
    _cover(document)

    lines = source.read_text(encoding="utf-8").splitlines()
    start = next(
        index for index, value in enumerate(lines) if value.strip() == "## 摘要"
    )
    inserted_overview = False
    inserted_wrist = False
    previous_kind: str | None = None
    current_decimal_num_id = decimal_num_id
    for kind, value in _parse_blocks(lines[start:]):
        if kind == "h1":
            document.add_heading(str(value), level=1)
        elif kind == "h2":
            document.add_heading(str(value), level=2)
        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.widow_control = True
            run = paragraph.add_run(str(value))
            _set_run_font(run, size=11)
            if (
                not inserted_overview
                and str(value).startswith("field-v3 来自用户提供的")
            ):
                _add_figure(
                    document,
                    FIELD_OVERVIEW,
                    "图 1  field-v3 草莓田场景总览",
                    6.15,
                )
                inserted_overview = True
            if (
                not inserted_wrist
                and str(value).startswith("固定目标的无运动检查")
            ):
                _add_figure(
                    document,
                    WRIST_IMAGE,
                    "图 2  腕部 RGB-D 相机的 field-v3 目标视图",
                    5.6,
                )
                inserted_wrist = True
        elif kind in {"bullet", "number"}:
            if kind == "number" and previous_kind != "number":
                current_decimal_num_id = _clone_numbering_instance(
                    document, decimal_num_id
                )
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.208
            _set_list_numbering(
                paragraph,
                bullet_num_id if kind == "bullet" else current_decimal_num_id,
            )
            run = paragraph.add_run(str(value))
            _set_run_font(run, size=11)
        elif kind == "code":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.right_indent = Inches(0.18)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.line_spacing = 1.05
            p_pr = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F2F4F7")
            shading.set(qn("w:val"), "clear")
            p_pr.insert(0, shading)
            run = paragraph.add_run(str(value))
            _set_run_font(
                run,
                font="Courier New",
                latin_font="Courier New",
                size=9.5,
                color="222222",
            )
        elif kind == "table":
            rows = value
            if not rows:
                continue
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.rows[0]._tr.get_or_add_trPr().append(
                OxmlElement("w:tblHeader")
            )
            for row_index, row in enumerate(rows):
                for column_index, cell_value in enumerate(row):
                    cell = table.cell(row_index, column_index)
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.15
                    if column_index > 0 and len(cell_value) < 28:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = paragraph.add_run(cell_value)
                    _set_run_font(
                        run,
                        size=9.5,
                        bold=row_index == 0,
                        color="1A1A1A",
                    )
                    if row_index == 0:
                        _shade_cell(cell, TABLE_FILL)
                    elif row_index % 2 == 0:
                        _shade_cell(cell, TABLE_ALT)
            _table_geometry(table, _column_widths(rows[0]))
            after = document.add_paragraph()
            after.paragraph_format.space_before = Pt(0)
            after.paragraph_format.space_after = Pt(0)
        previous_kind = kind

    output.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = (
        "基于 ROS 2、YOLO 与 Gazebo 的草莓成熟度识别及机械臂抓取搬运仿真"
    )
    document.core_properties.subject = "本科生 URP 提交版项目总结报告"
    document.core_properties.author = "草莓 URP 项目组"
    document.core_properties.comments = (
        "Preset: narrative_proposal; header: editorial_cover; "
        "named override: Chinese academic typography."
    )
    document.save(output)
    receipt = {
        "schema_version": 1,
        "kind": "submission_report_build",
        "preset": "narrative_proposal",
        "header_pattern": "editorial_cover",
        "named_override": "Chinese academic typography",
        "source": {
            "path": source.relative_to(REPOSITORY_ROOT).as_posix(),
            "sha256": _sha256(source),
        },
        "figures": [
            {
                "path": path.relative_to(REPOSITORY_ROOT).as_posix(),
                "sha256": _sha256(path),
            }
            for path in (FIELD_OVERVIEW, WRIST_IMAGE)
        ],
        "docx": {
            "path": output.relative_to(REPOSITORY_ROOT).as_posix(),
            "size_bytes": output.stat().st_size,
            "sha256": _sha256(output),
        },
    }
    receipt_path = output.with_suffix(".build-receipt.json")
    receipt_path.write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return receipt


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    arguments = parser.parse_args()
    receipt = build(arguments.source.resolve(), arguments.output.resolve())
    print(json.dumps(receipt, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
